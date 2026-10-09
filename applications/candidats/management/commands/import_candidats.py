import csv
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from applications.candidats.models import CANDIDATE_CODE_PATTERN, Candidat, NoteMatiere

REQUIRED_COLUMNS = {"annee", "code", "nom", "prenom"}
OPTIONAL_COLUMNS = {
    "date_naissance",
    "lieu_naissance",
    "telephone",
    "total_points",
    "total_sur",
    "mention",
    "type_resultat",
    "matiere",
    "note",
    "coefficient",
    "serie",
    "ecole",
}


def _parse_date(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    value = str(value).strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(value, pattern).date()
        except ValueError:
            continue
    raise ValueError(f"Date de naissance invalide : {value}")


def _decimal(value, default="0"):
    if value in (None, ""):
        return Decimal(default)
    try:
        return Decimal(str(value).strip().replace(",", "."))
    except InvalidOperation as exc:
        raise ValueError(f"Nombre invalide : {value}") from exc


def _rows(path):
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as file:
            yield from csv.DictReader(file)
        return
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise CommandError("Format non supporté. Utilisez un fichier CSV ou XLSX.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        values = workbook.active.iter_rows(values_only=True)
        headers = next(values, None)
        if headers:
            for row in values:
                yield dict(zip(headers, row))
    finally:
        workbook.close()


class Command(BaseCommand):
    help = "Importe ou met à jour les candidats à partir d'un CSV ou XLSX."

    def add_arguments(self, parser):
        parser.add_argument("fichier", type=Path)

    def handle(self, *args, **options):
        path = options["fichier"]
        if not path.is_file():
            raise CommandError(f"Fichier introuvable : {path}")
        imported = updated = 0
        for row_number, raw_row in enumerate(_rows(path), start=2):
            row = {
                str(key).strip().lower(): value
                for key, value in raw_row.items()
                if key is not None
            }
            missing = REQUIRED_COLUMNS - row.keys()
            if missing:
                raise CommandError(
                    f"Colonnes obligatoires manquantes à la ligne {row_number} : {', '.join(sorted(missing))}"
                )
            values = {key: row.get(key) for key in REQUIRED_COLUMNS | OPTIONAL_COLUMNS}
            try:
                values["annee"] = int(values["annee"])
                values["code"] = str(values["code"]).strip()
                if not re.fullmatch(CANDIDATE_CODE_PATTERN, values["code"]):
                    raise ValueError(
                        "Le matricule doit contenir 10 chiffres et commencer par 01 à 10."
                    )
                values["nom"] = str(values["nom"]).strip().upper()
                values["prenom"] = str(values["prenom"]).strip()
                values["date_naissance"] = _parse_date(values["date_naissance"])
                values["total_points"] = _decimal(values["total_points"])
                values["total_sur"] = int(values["total_sur"] or 1900)
                values["type_resultat"] = (
                    str(values["type_resultat"] or Candidat.TypeResultat.BAC_UNIQUE)
                    .strip()
                    .lower()
                )
                if values["type_resultat"] not in Candidat.TypeResultat.values:
                    raise ValueError(
                        f"Type de résultat invalide : {values['type_resultat']}"
                    )
                subject_values = None
                subject_fields = ("matiere", "note", "coefficient")
                has_subject = any(
                    values[field] not in (None, "") for field in subject_fields
                )
                if has_subject:
                    if not all(
                        values[field] not in (None, "") for field in subject_fields
                    ):
                        raise ValueError(
                            "Une matière doit avoir une note et un coefficient."
                        )
                    subject_values = {
                        "matiere": str(values["matiere"]).strip(),
                        "note": int(values["note"]),
                        "coefficient": int(values["coefficient"]),
                    }
                    if (
                        not subject_values["matiere"]
                        or subject_values["coefficient"] <= 0
                        or subject_values["note"] > subject_values["coefficient"]
                    ):
                        raise ValueError("Note ou coefficient de matière invalide.")
                for key in OPTIONAL_COLUMNS - {
                    "date_naissance",
                    "total_points",
                    "total_sur",
                    "type_resultat",
                    "matiere",
                    "note",
                    "coefficient",
                }:
                    values[key] = str(values[key] or "").strip()
            except (TypeError, ValueError, InvalidOperation) as exc:
                raise CommandError(f"Ligne {row_number} invalide : {exc}") from exc
            for field in ("matiere", "note", "coefficient"):
                values.pop(field, None)
            candidat, created = Candidat.objects.update_or_create(
                annee=values.pop("annee"), code=values.pop("code"), defaults=values
            )
            if subject_values:
                NoteMatiere.objects.update_or_create(
                    candidat=candidat,
                    matiere=subject_values["matiere"],
                    defaults={
                        "note": subject_values["note"],
                        "coefficient": subject_values["coefficient"],
                    },
                )
            imported += int(created)
            updated += int(not created)
        self.stdout.write(
            self.style.SUCCESS(
                f"Import terminé : {imported} créé(s), {updated} mis à jour."
            )
        )
